"""
Download and extract content from WhatsApp voice notes and document attachments.
"""

from __future__ import annotations

import base64
import datetime
import json
import os
import re
import shutil
import subprocess
import time
import requests
from typing import Any, Dict, List, Optional, Tuple

from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.keys import Keys
from openai import OpenAI

BASE_DIR = "/Users/evon/OpenClaw"
WA_AUDIO_DIR = os.path.join(BASE_DIR, "WA_Audio")
WA_IMAGE_DIR = os.path.join(BASE_DIR, "WA_Image")
WA_FILES_DIR = os.path.join(BASE_DIR, "WA_Files")
VOICE_CAPTURE_DIR = WA_AUDIO_DIR
VOICE_LATEST_OPUS = os.path.join(WA_AUDIO_DIR, "latest.opus")
VOICE_LATEST_WAV = os.path.join(WA_AUDIO_DIR, "latest.whisper.wav")
IMAGE_LATEST_PATH = os.path.join(WA_IMAGE_DIR, "latest.png")
DOC_CAPTURE_DIR = WA_FILES_DIR
DOC_PREVIEW_DIR = WA_IMAGE_DIR

VERSION = "v1.22-WA-IMAGE-VALIDATION"

COPILOT_BASE_URL = os.getenv("COPILOT_BASE_URL", "http://127.0.0.1:8000/v1")
COPILOT_MODEL = os.getenv("COPILOT_MODEL", "copilot")
WHISPER_MODEL = os.getenv("OPENCLAW_WHISPER_MODEL", "small")
WHISPER_LANGUAGE = os.getenv("OPENCLAW_WHISPER_LANGUAGE", "auto")
WHISPER_NORMALIZE = os.getenv("OPENCLAW_WHISPER_NORMALIZE", "0").strip().lower() in ("1", "true", "yes")
WHISPER_PROMPT = os.getenv(
    "OPENCLAW_WHISPER_PROMPT",
    "Quote quotation RFQ part number qty pieces meter E3Z-T61 报价 berapa harga pcs unit industrial parts.",
)
COPILOT_WHISPER_MODEL = os.getenv("COPILOT_WHISPER_MODEL", "whisper-1")

BLOB_TO_BASE64_JS = """
var el = arguments[0];
var callback = arguments[arguments.length - 1];
var src = el.currentSrc || el.src || '';
if (!src) {
  callback(null);
} else if (src.indexOf('data:') === 0) {
  callback(src.split(',')[1]);
} else {
  fetch(src).then(function(response) {
    return response.blob();
  }).then(function(blob) {
    var reader = new FileReader();
    reader.onloadend = function() {
      var result = reader.result || '';
      callback(typeof result === 'string' ? result.split(',')[1] : null);
    };
    reader.onerror = function() { callback(null); };
    reader.readAsDataURL(blob);
  }).catch(function() { callback(null); });
}
"""

DOWNLOAD_VOICE_NO_PLAY_JS = """
var dataId = arguments[0];
var container = arguments[1];
var callback = arguments[arguments.length - 1];

function sleep(ms) {
  return new Promise(function(resolve) { setTimeout(resolve, ms); });
}

function dataIdFromContainer(root) {
  if (!root) return '';
  var id = root.getAttribute && root.getAttribute('data-id') || '';
  if (id) return id;
  var child = root.querySelector && root.querySelector('[data-id]');
  return child ? (child.getAttribute('data-id') || '') : '';
}

function toB64FromBytes(data) {
  var blob = new Blob([data], {type: 'audio/ogg; codecs=opus'});
  return new Promise(function(resolve, reject) {
    var reader = new FileReader();
    reader.onloadend = function() {
      var result = reader.result || '';
      resolve(typeof result === 'string' ? result.split(',')[1] : null);
    };
    reader.onerror = reject;
    reader.readAsDataURL(blob);
  });
}

function ensureStore() {
  if (window.__openclawStoreReady && window.Store && window.Store.Msg) return true;
  try {
    if (typeof window.require !== 'function' && window.webpackChunkwhatsapp_web_client) {
      window.webpackChunkwhatsapp_web_client.push([
        ['openclaw_store_boot'], {}, function(r) { window.require = r; }
      ]);
    }
  } catch (e0) {}
  if (typeof window.require !== 'function') return false;

  window.Store = window.Store || {};
  if (window.Store.Msg && window.Store.DownloadManager) {
    window.__openclawStoreReady = true;
    return true;
  }

  var pairs = [
    ['WAWebCollections', 'Msg'],
    ['WAWebCollections', 'WAWebDownloadManager'],
    ['WAWebMsgCollection', 'WAWebDownloadManager'],
    ['WAWebFrontendMsgCollection', 'WAWebDownloadManager'],
  ];
  for (var p = 0; p < pairs.length; p++) {
    try {
      var cols = window.require(pairs[p][0]);
      var dmMod = window.require(pairs[p][1]);
      window.Store.Msg = cols.Msg || (cols.default && cols.default.Msg) || cols;
      var dm = dmMod.downloadManager || (dmMod.default && dmMod.default.downloadManager) || dmMod;
      window.Store.DownloadManager = dm;
    } catch (e1) {}
  }

  try {
    var req = window.require;
    if (req.c) {
      for (var mid in req.c) {
        if (!req.c[mid] || !req.c[mid].exports) continue;
        var ex = req.c[mid].exports;
        var mods = [ex, ex.default];
        for (var mi = 0; mi < mods.length; mi++) {
          var m = mods[mi];
          if (!m || typeof m !== 'object') continue;
          if (m.Msg && typeof m.Msg.get === 'function' && !window.Store.Msg) {
            window.Store.Msg = m.Msg;
          }
          if (m.downloadManager && !window.Store.DownloadManager) {
            window.Store.DownloadManager = m.downloadManager;
          }
          if (typeof m.downloadAndDecrypt === 'function' && !window.Store.DownloadManager) {
            window.Store.DownloadManager = m;
          }
        }
      }
    }
  } catch (e2) {}

  if (window.Store.Msg && window.Store.DownloadManager) {
    window.__openclawStoreReady = true;
    return true;
  }
  return !!window.Store.Msg;
}

function findMsg(id) {
  if (!id || !window.Store || !window.Store.Msg) return null;
  try {
    var direct = window.Store.Msg.get(id);
    if (direct) return direct;
  } catch (e) {}
  var list = [];
  try {
    if (window.Store.Msg._models) list = window.Store.Msg._models;
    else if (window.Store.Msg.models) list = window.Store.Msg.models;
    else if (typeof window.Store.Msg.getModelsArray === 'function') {
      list = window.Store.Msg.getModelsArray();
    }
  } catch (e2) { list = []; }
  if (typeof list.forEach !== 'function' && list.length === undefined) {
    try { list = Array.from(list); } catch (e3) { list = []; }
  }
  var suffix = String(id).slice(-20);
  for (var i = 0; i < list.length; i++) {
    var m = list[i];
    var sid = '';
    try {
      sid = m.id && (m.id._serialized || m.id.id || String(m.id)) || '';
    } catch (e4) { sid = ''; }
    if (!sid) continue;
    if (sid === id || sid.indexOf(id) >= 0 || id.indexOf(sid) >= 0) return m;
    if (sid.slice(-20) === suffix || sid.endsWith(id) || id.endsWith(sid.slice(-16))) return m;
  }
  return null;
}

function downloadViaManager(msg) {
  var dm = window.Store && window.Store.DownloadManager;
  if (!dm) return Promise.resolve(null);
  var md = msg.mediaData || msg;
  if (typeof msg.downloadMedia === 'function') {
    return msg.downloadMedia().then(function(data) {
      if (!data) return null;
      if (data instanceof ArrayBuffer) return toB64FromBytes(data);
      if (data.buffer) return toB64FromBytes(data.buffer);
      if (data.arrayBuffer) return data.arrayBuffer().then(toB64FromBytes);
      return toB64FromBytes(data);
    }).catch(function() { return null; });
  }
  if (typeof dm.downloadAndDecrypt !== 'function') {
    return Promise.resolve(null);
  }
  return dm.downloadAndDecrypt({
    directPath: md.directPath,
    encFilehash: md.encFilehash,
    filehash: md.filehash,
    mediaKey: md.mediaKey,
    mediaKeyTimestamp: md.mediaKeyTimestamp,
    type: md.type || msg.type || 'ptt',
    signal: (new AbortController()).signal
  }).then(function(data) {
    if (!data) return null;
    if (data instanceof ArrayBuffer) return toB64FromBytes(data);
    if (data.buffer) return toB64FromBytes(data.buffer);
    return toB64FromBytes(data);
  }).catch(function() { return null; });
}

(function run() {
  try {
    var closer = document.querySelector('[data-testid="btn-closer-drawer"], [data-testid="drawer-back"]');
    if (closer) { try { closer.click(); } catch (e) {} }
    if (container) {
      try { container.scrollIntoView({block: 'center', inline: 'center'}); } catch (e2) {}
    }
    var id = String(dataId || dataIdFromContainer(container) || '');
    if (!id || !ensureStore()) {
      callback(null);
      return;
    }
    var msg = findMsg(id);
    if (!msg) {
      callback(null);
      return;
    }
    downloadViaManager(msg).then(function(b64) {
      callback(b64 || null);
    }).catch(function() { callback(null); });
  } catch (e) {
    callback(null);
  }
})();
"""

CAPTURE_HOOKED_BLOB_JS = """
var container = arguments[0];
var callback = arguments[arguments.length - 1];

function sleep(ms) {
  return new Promise(function(resolve) { setTimeout(resolve, ms); });
}

function muteAll() {
  document.querySelectorAll('audio, video').forEach(function(el) {
    try { el.muted = true; el.volume = 0; } catch (e) {}
  });
}

function clickPlay(root) {
  if (!root) return false;
  var selectors = [
    '[data-testid="audio-play"]',
    '[data-testid="ptt-play-button"]',
    '[data-testid="ptt"]',
    'span[data-icon="audio-play"]',
    'span[data-icon="ptt"]',
    '[role="button"][aria-label*="Play" i]',
    'canvas'
  ];
  for (var s = 0; s < selectors.length; s++) {
    var btn = root.querySelector(selectors[s]);
    if (btn) {
      try {
        btn.click();
        return true;
      } catch (e) {}
    }
  }
  try {
    root.click();
    return true;
  } catch (e2) {}
  return false;
}

function isOggBuffer(buf) {
  var b = new Uint8Array(buf);
  return b.length >= 128 &&
    b[0] === 0x4F && b[1] === 0x67 && b[2] === 0x67 && b[3] === 0x53;
}

function bufferToB64(buf) {
  var b = new Uint8Array(buf);
  var binary = '';
  var chunk = 0x8000;
  for (var i = 0; i < b.length; i += chunk) {
    binary += String.fromCharCode.apply(null, b.subarray(i, i + chunk));
  }
  return btoa(binary);
}

function fetchOggB64(url) {
  if (!url) return Promise.resolve(null);
  return fetch(url).then(function(r) { return r.arrayBuffer(); }).then(function(buf) {
    if (!isOggBuffer(buf)) return null;
    return bufferToB64(buf);
  }).catch(function() { return null; });
}

function tryUrls(urls, idx) {
  if (idx >= urls.length) return Promise.resolve(null);
  return fetchOggB64(urls[idx]).then(function(b64) {
    if (b64) return b64;
    return tryUrls(urls, idx + 1);
  });
}

function isMediaUrl(url) {
  if (!url || typeof url !== 'string') return false;
  return url.indexOf('blob:') === 0 || url.indexOf('mmg.whatsapp.net') > -1;
}

function freshBlobUrls(startIdx) {
  var urls = window.__openclawVoiceUrls || [];
  var out = [];
  for (var i = startIdx; i < urls.length; i++) {
    var u = urls[i];
    if (isMediaUrl(u) && out.indexOf(u) < 0) out.push(u);
  }
  if (container) {
    container.querySelectorAll('audio').forEach(function(a) {
      var u = a.currentSrc || a.src || '';
      if (isMediaUrl(u) && out.indexOf(u) < 0) out.push(u);
    });
  }
  return out;
}

function poll(startIdx, deadline) {
  var urls = freshBlobUrls(startIdx);
  if (urls.length) {
    return tryUrls(urls, 0).then(function(b64) {
      if (b64) return b64;
      if (Date.now() < deadline) return sleep(500).then(function() { return poll(startIdx, deadline); });
      return null;
    });
  }
  if (Date.now() < deadline) return sleep(500).then(function() { return poll(startIdx, deadline); });
  return Promise.resolve(null);
}

(function run() {
  window.__openclawVoiceUrls = window.__openclawVoiceUrls || [];
  var startIdx = window.__openclawVoiceUrls.length;
  var closer = document.querySelector('[data-testid="btn-closer-drawer"], [data-testid="drawer-back"]');
  if (closer) { try { closer.click(); } catch (e) {} }
  if (container) {
    try { container.scrollIntoView({block: 'center', inline: 'center'}); } catch (e2) {}
  }
  muteAll();
  clickPlay(container);
  sleep(1000).then(function() {
    return poll(startIdx, Date.now() + 22000);
  }).then(function(b64) {
    callback(b64 || null);
  });
})();
"""

CAPTURE_VOICE_OGG_JS = """
var container = arguments[0];
var callback = arguments[arguments.length - 1];

function sleep(ms) {
  return new Promise(function(resolve) { setTimeout(resolve, ms); });
}

function snapshotAudioUrls() {
  var urls = [];
  document.querySelectorAll('audio').forEach(function(a) {
    var u = a.currentSrc || a.src || '';
    if (u && urls.indexOf(u) < 0) urls.push(u);
  });
  return urls;
}

function muteAll() {
  document.querySelectorAll('audio, video').forEach(function(el) {
    try { el.muted = true; el.volume = 0; } catch (e) {}
  });
}

function clickPlay(root) {
  if (!root) return false;
  var selectors = [
    '[data-testid="audio-play"]',
    '[data-testid="ptt-play-button"]',
    'span[data-icon="audio-play"]',
    'span[data-icon="ptt"]',
    '[role="button"][aria-label*="Play" i]'
  ];
  for (var s = 0; s < selectors.length; s++) {
    var btn = root.querySelector(selectors[s]);
    if (btn) { try { btn.click(); return true; } catch (e) {} }
  }
  return false;
}

function isOggBuffer(buf) {
  var b = new Uint8Array(buf);
  return b.length >= 128 &&
    b[0] === 0x4F && b[1] === 0x67 && b[2] === 0x67 && b[3] === 0x53;
}

function bufferToB64(buf) {
  var b = new Uint8Array(buf);
  var binary = '';
  var chunk = 0x8000;
  for (var i = 0; i < b.length; i += chunk) {
    binary += String.fromCharCode.apply(null, b.subarray(i, i + chunk));
  }
  return btoa(binary);
}

function fetchOggB64(url) {
  if (!url) return Promise.resolve(null);
  if (url.indexOf('blob:') !== 0 && url.indexOf('mmg.whatsapp.net') < 0) return Promise.resolve(null);
  return fetch(url).then(function(r) { return r.arrayBuffer(); }).then(function(buf) {
    if (!isOggBuffer(buf)) return null;
    return bufferToB64(buf);
  }).catch(function() { return null; });
}

function newAudioUrls(before) {
  var out = [];
  if (container) {
    container.querySelectorAll('audio').forEach(function(a) {
      var u = a.currentSrc || a.src || '';
      if (u && before.indexOf(u) < 0 && out.indexOf(u) < 0) out.push(u);
    });
  }
  document.querySelectorAll('#main audio, audio').forEach(function(a) {
    var u = a.currentSrc || a.src || '';
    if (u && before.indexOf(u) < 0 && out.indexOf(u) < 0) out.push(u);
  });
  return out;
}

function tryUrls(urls, idx) {
  if (idx >= urls.length) return Promise.resolve(null);
  return fetchOggB64(urls[idx]).then(function(b64) {
    if (b64) return b64;
    return tryUrls(urls, idx + 1);
  });
}

function poll(before, deadline) {
  var urls = newAudioUrls(before);
  if (urls.length) {
    return tryUrls(urls, 0).then(function(b64) {
      if (b64) return b64;
      if (Date.now() < deadline) return sleep(400).then(function() { return poll(before, deadline); });
      return null;
    });
  }
  if (Date.now() < deadline) return sleep(400).then(function() { return poll(before, deadline); });
  return Promise.resolve(null);
}

(function run() {
  var closer = document.querySelector('[data-testid="btn-closer-drawer"], [data-testid="drawer-back"]');
  if (closer) { try { closer.click(); } catch (e) {} }
  if (container) {
    try { container.scrollIntoView({block: 'center', inline: 'center'}); } catch (e2) {}
  }
  var before = snapshotAudioUrls();
  muteAll();
  clickPlay(container);
  sleep(900).then(function() {
    return poll(before, Date.now() + 18000);
  }).then(function(b64) {
    callback(b64 || null);
  });
})();
"""

RELOCATE_LAST_VOICE_JS = """
const panel = document.querySelector('#main [data-testid="conversation-panel-messages"]')
    || document.querySelector('#main');
if (!panel) return null;

function isOutgoing(container) {
    let node = container;
    for (let depth = 0; depth < 8 && node; depth++) {
        if (node.classList && node.classList.contains('message-out')) return true;
        if (node.classList && node.classList.contains('message-in')) return false;
        node = node.parentElement;
    }
    return false;
}

function hasVoice(container) {
    if (container.querySelector(
        'audio, [data-testid="audio-play"], [data-testid="ptt-play-button"], [data-testid="ptt"], '
        + '[data-testid="audio"], [data-icon="ptt"], [data-icon="audio-play"]'
    )) return true;
    const inner = container.innerText || '';
    const hasDuration = /\\b\\d{1,2}:\\d{2}\\b/.test(inner);
    const hasPlay = !!container.querySelector(
        'span[data-icon="audio-play"], span[data-icon="ptt"], [data-testid="audio-play"], canvas'
    );
    return hasDuration && hasPlay;
}

const voices = [];
const seen = new Set();
for (const node of panel.querySelectorAll('[data-testid="msg-container"], div.message-in')) {
    const container = node.matches('[data-testid="msg-container"]')
        ? node
        : (node.closest('[data-testid="msg-container"]') || node);
    if (!container || seen.has(container)) continue;
    if (isOutgoing(container)) continue;
    if (!hasVoice(container)) continue;
    seen.add(container);
    voices.push(container);
}
return voices.length ? voices[voices.length - 1] : null;
"""

RELOCATE_MSG_CONTAINER_JS = """
var dataId = arguments[0];
if (!dataId) return null;
var panel = document.querySelector('#main');
if (!panel) return null;
var nodes = panel.querySelectorAll('[data-id]');
for (var i = 0; i < nodes.length; i++) {
  var id = nodes[i].getAttribute('data-id') || '';
  if (id === dataId || id.indexOf(dataId) === 0) {
    var container = nodes[i].closest('[data-testid="msg-container"]');
    return container || nodes[i];
  }
}
return null;
"""

PLAY_AND_CAPTURE_VOICE_JS = DOWNLOAD_VOICE_NO_PLAY_JS


def is_valid_voice_bytes(raw: bytes) -> bool:
    """True when bytes look like a real opus/wav file (same quality bar as manual curl upload)."""
    if len(raw) < 128:
        return False
    if raw.startswith(b"OggS"):
        return True
    if raw.startswith(b"RIFF") and len(raw) > 12 and raw[8:12] == b"WAVE":
        return True
    return False


def _safe_contact(contact_name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(contact_name or "contact"))[:60]


def _timestamp_slug() -> str:
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def _find_voice_element(bubble):
    if bubble is None:
        return None
    selectors = [
        "audio",
        '[data-testid="audio-play"]',
        '[data-testid="ptt-play-button"]',
        '[data-testid="ptt"]',
        '[data-testid="audio"]',
        '[data-icon="audio-play"]',
        '[data-icon="ptt"]',
        "canvas",
    ]
    for selector in selectors:
        try:
            elements = bubble.find_elements(By.CSS_SELECTOR, selector)
            if elements:
                for element in elements:
                    if element.tag_name.lower() == "audio":
                        return element
                    parent = element
                    for _ in range(4):
                        try:
                            audio = parent.find_element(By.CSS_SELECTOR, "audio")
                            if audio:
                                return audio
                        except Exception:
                            pass
                        try:
                            parent = parent.find_element(By.XPATH, "..")
                        except Exception:
                            break
                return elements[0]
        except Exception:
            continue
    return None


def _find_document_element(bubble):
    if bubble is None:
        return None
    selectors = [
        '[data-testid="document-thumb"]',
        '[data-testid="document"]',
        '[data-testid="document-message"]',
        '[data-icon="document-pdf"]',
        '[data-icon="document"]',
        '[data-icon="document-xls"]',
        '[data-icon="document-doc"]',
        'span[data-testid="document-thumb"]',
        'div[role="button"][aria-label*="pdf" i]',
        'div[role="button"][aria-label*="document" i]',
    ]
    for selector in selectors:
        try:
            elements = bubble.find_elements(By.CSS_SELECTOR, selector)
            if elements:
                return elements[0]
        except Exception:
            continue
    try:
        if re.search(r"\.pdf\b", bubble.text or "", re.I):
            buttons = bubble.find_elements(By.CSS_SELECTOR, 'div[role="button"]')
            if buttons:
                return buttons[0]
    except Exception:
        pass
    return None


def _find_download_in_bubble(bubble):
    if bubble is None:
        return None
    selectors = [
        '[data-testid="download"]',
        'span[data-icon="download"]',
        '[data-icon="download"]',
        '[aria-label="Download"]',
        '[title="Download"]',
    ]
    for selector in selectors:
        try:
            for element in bubble.find_elements(By.CSS_SELECTOR, selector):
                if element.is_displayed():
                    return element
        except Exception:
            continue
    return None


def capture_document_preview_image(driver, bubble, contact_name: str) -> Optional[str]:
    """Screenshot PDF/document preview inside bubble for Copilot vision fallback."""
    if bubble is None:
        return None
    os.makedirs(DOC_PREVIEW_DIR, exist_ok=True)
    out_path = os.path.join(
        DOC_PREVIEW_DIR,
        f"{_timestamp_slug()}_{_safe_contact(contact_name)}_preview.png",
    )
    try:
        driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center', inline: 'center'});",
            bubble,
        )
        time.sleep(0.5)
        bubble.screenshot(out_path)
        print(f"📄 [DOC] Saved document preview screenshot → {WA_IMAGE_DIR}/{os.path.basename(out_path)}")
        return out_path
    except Exception as exc:
        print(f"⚠️ [DOC] Preview screenshot failed: {exc}")
        return None


def _click_voice_play(driver, bubble) -> bool:
    """Click play on a PTT bubble so WhatsApp loads the audio element."""
    clicked = False
    for selector in (
        '[data-testid="audio-play"]',
        '[data-testid="ptt-play-button"]',
        'span[data-icon="audio-play"]',
        'span[data-icon="ptt"]',
        '[role="button"][aria-label*="Play" i]',
        'button[aria-label*="Play" i]',
    ):
        try:
            for element in bubble.find_elements(By.CSS_SELECTOR, selector):
                if element.is_displayed():
                    driver.execute_script("arguments[0].click();", element)
                    clicked = True
                    time.sleep(1.0)
        except Exception:
            continue
    if not clicked:
        try:
            driver.execute_script("arguments[0].click();", bubble)
            clicked = True
            time.sleep(1.0)
        except Exception:
            pass
    if clicked:
        time.sleep(2.0)
    return clicked


def _resolve_voice_container(bubble):
    if bubble is None:
        return None
    try:
        testid = bubble.get_attribute("data-testid") or ""
        if testid == "msg-container" or "message-in" in (bubble.get_attribute("class") or ""):
            return bubble
        for _ in range(8):
            bubble = bubble.find_element(By.XPATH, "..")
            testid = bubble.get_attribute("data-testid") or ""
            if testid == "msg-container" or "message-in" in (bubble.get_attribute("class") or ""):
                return bubble
    except Exception:
        pass
    return bubble


def _guess_voice_extension(data: bytes) -> str:
    if data.startswith(b"OggS"):
        return ".opus"
    if data.startswith(b"RIFF"):
        return ".wav"
    if data.startswith(b"ID3") or (len(data) > 2 and data[0] == 0xFF and (data[1] & 0xE0) == 0xE0):
        return ".mp3"
    return ".opus"


def convert_voice_to_wav(audio_path: str, out_wav: str | None = None) -> tuple[str, bool]:
    """Convert Ogg/Opus to WAV. out_wav overwrites each run (e.g. latest.whisper.wav)."""
    if audio_path.lower().endswith(".wav"):
        return audio_path, False

    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        print("⚠️ [VOICE] ffmpeg not found — install: brew install ffmpeg")
        return audio_path, False

    wav_path = out_wav or f"{audio_path}.whisper.wav"
    print(f"🎤 [VOICE] ffmpeg {os.path.basename(audio_path)} → {os.path.basename(wav_path)} (overwrite)")
    result = subprocess.run(
        [
            ffmpeg, "-nostdin", "-hide_banner", "-loglevel", "error", "-y",
            "-i", audio_path, "-ar", "16000", "-ac", "1", "-c:a", "pcm_s16le",
            wav_path,
        ],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "").strip()
        print(f"⚠️ [VOICE] ffmpeg convert failed: {err[-300:]}")
        return audio_path, False

    if os.path.exists(wav_path) and os.path.getsize(wav_path) > 44:
        print(f"🎤 [VOICE] ffmpeg ok: {os.path.getsize(wav_path)} bytes → {wav_path}")
        return wav_path, False
    return audio_path, False


def _execute_async_js(driver, script, *args, timeout: int = 60):
    """Run execute_async_script with a generous timeout."""
    try:
        driver.set_script_timeout(timeout)
        return driver.execute_async_script(script, *args)
    except Exception:
        raise


def _resolve_data_id(container, message_data_id: str = "") -> str:
    data_id = str(message_data_id or "").strip()
    if data_id:
        return data_id
    if container is None:
        return ""
    try:
        data_id = str(container.get_attribute("data-id") or "").strip()
        if data_id:
            return data_id
        child = container.find_elements(By.CSS_SELECTOR, "[data-id]")
        if child:
            data_id = str(child[0].get_attribute("data-id") or "").strip()
    except Exception:
        pass
    return data_id


def _relocate_voice_container(driver, bubble, message_data_id: str = ""):
    """Refresh the live voice bubble element (stale refs break blob capture)."""
    data_id = str(message_data_id or "").strip()
    if data_id:
        try:
            relocated = driver.execute_script(RELOCATE_MSG_CONTAINER_JS, data_id)
            if relocated is not None:
                print(f"🎤 [VOICE] Relocated bubble via data-id {data_id[:32]!r}")
                return relocated
        except Exception as exc:
            print(f"⚠️ [VOICE] data-id relocation failed: {exc}")

    resolved = _resolve_voice_container(bubble)
    if resolved is not None:
        return resolved

    try:
        relocated = driver.execute_script(RELOCATE_LAST_VOICE_JS)
        if relocated is not None:
            print("🎤 [VOICE] Relocated bubble via last-incoming-voice fallback")
            return relocated
    except Exception as exc:
        print(f"⚠️ [VOICE] last-incoming-voice fallback failed: {exc}")
    return None


def install_voice_capture_hooks(driver) -> None:
    """Ensure WhatsApp page hooks media URLs (re-run each download — page may have navigated)."""
    hook_js = """
window.__openclawVoiceUrls = window.__openclawVoiceUrls || [];
if (!window.__openclawVoiceHooked) {
  window.__openclawVoiceHooked = true;
  const push = (url) => {
    if (!url || typeof url !== 'string') return;
    if (url.indexOf('blob:') === 0 || url.indexOf('mmg.whatsapp.net') > -1) {
      window.__openclawVoiceUrls.push(url);
    }
  };
  const origFetch = window.fetch;
  window.fetch = function(input, init) {
    push(typeof input === 'string' ? input : (input && input.url) || '');
    return origFetch.apply(this, arguments);
  };
  const origOpen = XMLHttpRequest.prototype.open;
  XMLHttpRequest.prototype.open = function(method, url) {
    push(String(url || ''));
    return origOpen.apply(this, arguments);
  };
  if (window.URL && window.URL.createObjectURL) {
    const origCreate = URL.createObjectURL;
    URL.createObjectURL = function(blob) {
      const url = origCreate.call(URL, blob);
      push(url);
      return url;
    };
  }
}
"""
    try:
        driver.execute_cdp_cmd(
            "Page.addScriptToEvaluateOnNewDocument",
            {"source": hook_js},
        )
    except Exception:
        pass
    try:
        driver.execute_script(hook_js)
        driver._openclaw_voice_hooks_installed = True
    except Exception:
        pass


def _save_voice_to_wa_audio(raw: bytes, message_data_id: str = "") -> Optional[str]:
    """Save valid opus to WA_Audio/{data_id}_{timestamp}.opus + latest.opus."""
    if not is_valid_voice_bytes(raw):
        header = raw[:16].hex() if raw else "empty"
        print(f"⚠️ [VOICE] Not opus audio ({len(raw)} bytes, header={header}) — skip save.")
        return None

    os.makedirs(WA_AUDIO_DIR, exist_ok=True)
    safe_id = re.sub(r"[^A-Za-z0-9._-]+", "_", message_data_id or "voice")[:48]
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    named_path = os.path.join(WA_AUDIO_DIR, f"{safe_id}_{stamp}.opus")

    with open(named_path, "wb") as f:
        f.write(raw)
    with open(VOICE_LATEST_OPUS, "wb") as f:
        f.write(raw)

    manifest_path = os.path.join(WA_AUDIO_DIR, "latest.json")
    try:
        with open(manifest_path, "w", encoding="utf-8") as f:
            json.dump(
                {
                    "saved_at": datetime.datetime.now().isoformat(timespec="seconds"),
                    "data_id": message_data_id,
                    "named_path": named_path,
                    "bytes": len(raw),
                },
                f,
                indent=2,
            )
    except Exception:
        pass

    print(f"✅ [VOICE] Downloaded → {named_path} ({len(raw)} bytes)")
    print(f"✅ [VOICE] Updated {VOICE_LATEST_OPUS}")

    wav_path, _ = convert_voice_to_wav(VOICE_LATEST_OPUS, out_wav=VOICE_LATEST_WAV)
    if wav_path.lower().endswith(".wav") and os.path.exists(wav_path) and os.path.getsize(wav_path) > 44:
        print(f"🎤 [VOICE] ffmpeg sidecar: {wav_path}")
    # Return .opus path — matches manual curl: -F file=@.../latest.opus
    return named_path


def _try_capture_hooked_blob(driver, container) -> Optional[bytes]:
    """Muted play; capture Ogg Opus from createObjectURL / fetch hooks."""
    print("🎤 [VOICE] Trying download method: capture-hooked-blob (muted play)...")
    return _try_download_b64(
        driver,
        "capture-hooked-blob",
        CAPTURE_HOOKED_BLOB_JS,
        container,
        quiet=True,
        timeout=90,
    )


def _try_capture_ogg_audio(driver, container) -> Optional[bytes]:
    """Muted single play; fetch only NEW blob: URLs that decode as Ogg Opus."""
    print("🎤 [VOICE] Trying download method: capture-ogg-blob...")
    return _try_download_b64(
        driver, "capture-ogg-blob", CAPTURE_VOICE_OGG_JS, container, quiet=True, timeout=90
    )


def _try_download_b64(
    driver, method: str, script: str, *args, quiet: bool = False, timeout: int = 60
) -> Optional[bytes]:
    if not quiet:
        print(f"🎤 [VOICE] Trying download method: {method}...")
    try:
        b64 = _execute_async_js(driver, script, *args, timeout=timeout)
    except Exception as exc:
        print(f"⚠️ [VOICE] {method} failed: {exc}")
        return None
    if not b64:
        if not quiet:
            print(f"⚠️ [VOICE] {method} returned no data.")
        else:
            print(f"⚠️ [VOICE] {method} returned no data.")
        return None
    try:
        return base64.b64decode(b64)
    except Exception as exc:
        print(f"⚠️ [VOICE] {method} base64 decode failed: {exc}")
        return None


def _pick_newest_voice_download(min_mtime: float = 0) -> Optional[str]:
    """Find a freshly downloaded WhatsApp .opus file in ~/Downloads."""
    download_dirs = [
        os.path.expanduser("~/Downloads"),
        WA_AUDIO_DIR,
    ]
    candidates = []
    now = time.time()
    for directory in download_dirs:
        if not os.path.isdir(directory):
            continue
        for name in os.listdir(directory):
            path = os.path.join(directory, name)
            if not os.path.isfile(path):
                continue
            lower = name.lower()
            if not (
                lower.endswith((".opus", ".ogg", ".oga"))
                or ("whatsapp" in lower and "audio" in lower)
            ):
                continue
            try:
                mtime = os.path.getmtime(path)
            except OSError:
                continue
            if min_mtime and mtime <= min_mtime + 0.3:
                continue
            if now - mtime > 120:
                continue
            candidates.append((mtime, path))
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


OPEN_MESSAGE_MENU_JS = """
var container = arguments[0];
if (!container) return false;

function hoverEl(el) {
  if (!el) return;
  ['mouseover', 'mouseenter', 'mousemove'].forEach(function(type) {
    try {
      el.dispatchEvent(new MouseEvent(type, {bubbles: true, cancelable: true, view: window}));
    } catch (e) {}
  });
}

function scopes() {
  var out = [container];
  var mc = container.closest('[data-testid="msg-container"]');
  if (mc) {
    out.push(mc);
    if (mc.parentElement) out.push(mc.parentElement);
  }
  var row = container.closest('.message-in, .message-out, div[data-id]');
  if (row) out.push(row);
  return out.filter(Boolean);
}

var triggerSelectors = [
  '[data-icon="down"]',
  'span[data-icon="down"]',
  '[data-testid="down"]',
  '[aria-label*="message menu" i]',
  '[aria-label*="Open the message" i]',
  '[aria-label*="Menu" i]',
  'button[aria-label*="Menu" i]',
  '[role="button"][aria-label*="Menu" i]'
];

var seen = scopes();
for (var s = 0; s < seen.length; s++) {
  hoverEl(seen[s]);
}

for (var s = 0; s < seen.length; s++) {
  for (var i = 0; i < triggerSelectors.length; i++) {
    var el = seen[s].querySelector(triggerSelectors[i]);
    if (el) {
      try {
        el.click();
        return true;
      } catch (e) {}
    }
  }
}
return false;
"""

CLICK_MENU_DOWNLOAD_JS = """
function visible(el) {
  if (!el) return false;
  var r = el.getBoundingClientRect();
  return r.width > 0 && r.height > 0;
}

var selectors = [
  '[data-testid="mi-download"]',
  '[data-testid="message-download"]',
  '[data-icon="download"]',
  '[aria-label="Download"]',
  '[title="Download"]'
];

for (var i = 0; i < selectors.length; i++) {
  var nodes = document.querySelectorAll(selectors[i]);
  for (var j = 0; j < nodes.length; j++) {
    var el = nodes[j];
    if (!visible(el)) continue;
    var icon = el.getAttribute('data-icon') || '';
    if (icon === 'audio-download') continue;
    var parentText = (el.closest('[role="button"], [role="menuitem"], li, div') || el).innerText || '';
    if (parentText.indexOf('Download') >= 0 || icon === 'download' || el.getAttribute('data-testid')) {
      try {
        (el.closest('[role="button"], [role="menuitem"], li') || el).click();
        return true;
      } catch (e) {}
    }
  }
}

var buttons = document.querySelectorAll('[role="button"], [role="menuitem"], li');
for (var k = 0; k < buttons.length; k++) {
  var btn = buttons[k];
  if (!visible(btn)) continue;
  var text = (btn.innerText || btn.textContent || '').trim();
  if (text === 'Download') {
    try {
      btn.click();
      return true;
    } catch (e2) {}
  }
}
return false;
"""


def _message_menu_scopes(bubble):
    """Bubble + parent wrappers where WhatsApp hides the ▼ chevron."""
    scopes = [bubble]
    try:
        mc = bubble.find_element(
            By.XPATH,
            './ancestor-or-self::*[@data-testid="msg-container"][1]',
        )
        if mc not in scopes:
            scopes.append(mc)
        parent = mc.find_element(By.XPATH, "..")
        if parent not in scopes:
            scopes.append(parent)
    except Exception:
        pass
    try:
        row = bubble.find_element(
            By.XPATH,
            './ancestor::*[contains(@class,"message-in") or contains(@class,"message-out")][1]',
        )
        if row not in scopes:
            scopes.append(row)
    except Exception:
        pass
    return scopes


def _open_message_dropdown_menu(driver, bubble) -> bool:
    """Hover voice bubble and click ▼ to open Reply/React/Download menu."""
    if bubble is None:
        return False

    driver.execute_script(
        "arguments[0].scrollIntoView({block: 'center', inline: 'center'});",
        bubble,
    )
    time.sleep(0.5)

    try:
        ActionChains(driver).move_to_element(bubble).pause(0.8).perform()
    except Exception:
        pass

    trigger_selectors = [
        '[data-icon="down"]',
        'span[data-icon="down"]',
        '[data-testid="down"]',
        '[aria-label*="message menu" i]',
        '[aria-label*="Open the message" i]',
        'button[aria-label*="Menu" i]',
        '[role="button"][aria-label*="Menu" i]',
    ]

    for scope in _message_menu_scopes(bubble):
        try:
            ActionChains(driver).move_to_element(scope).pause(0.4).perform()
        except Exception:
            pass
        for selector in trigger_selectors:
            try:
                for el in scope.find_elements(By.CSS_SELECTOR, selector):
                    if not el.is_displayed():
                        continue
                    driver.execute_script("arguments[0].click();", el)
                    print(f"🎤 [VOICE] Opened message menu via {selector}")
                    time.sleep(0.7)
                    return True
            except Exception:
                continue

    try:
        opened = driver.execute_script(OPEN_MESSAGE_MENU_JS, bubble)
        if opened:
            print("🎤 [VOICE] Opened message menu via JS hover/click")
            time.sleep(0.7)
            return True
    except Exception:
        pass

    try:
        ActionChains(driver).context_click(bubble).perform()
        print("🎤 [VOICE] Opened message menu via right-click fallback")
        time.sleep(0.7)
        return True
    except Exception:
        return False


def _click_download_in_context_menu(driver) -> bool:
    """Click Download in the open message context menu."""
    time.sleep(0.5)

    xpath_selectors = [
        "//div[@role='button'][.//span[normalize-space()='Download']]",
        "//li[@role='button'][normalize-space()='Download']",
        "//*[@role='button' and normalize-space(.)='Download']",
        "//*[@role='menuitem' and normalize-space(.)='Download']",
    ]
    for xpath in xpath_selectors:
        try:
            for el in driver.find_elements(By.XPATH, xpath):
                if el.is_displayed():
                    driver.execute_script("arguments[0].click();", el)
                    print("🎤 [VOICE] Clicked Download in context menu")
                    return True
        except Exception:
            continue

    css_selectors = [
        '[data-testid="mi-download"]',
        '[data-testid="message-download"]',
        '[aria-label="Download"]',
    ]
    for selector in css_selectors:
        try:
            for el in driver.find_elements(By.CSS_SELECTOR, selector):
                if el.is_displayed():
                    driver.execute_script("arguments[0].click();", el)
                    print(f"🎤 [VOICE] Clicked Download via {selector}")
                    return True
        except Exception:
            continue

    try:
        if driver.execute_script(CLICK_MENU_DOWNLOAD_JS):
            print("🎤 [VOICE] Clicked Download via JS menu scan")
            return True
    except Exception:
        pass

    return False


def _try_download_voice_via_button(driver, bubble) -> Optional[bytes]:
    """Open message ▼ menu → Download; read the .opus saved to ~/Downloads."""
    print("🎤 [VOICE] Trying download method: message menu → Download...")
    if bubble is None:
        return None

    baseline = _pick_newest_voice_download()
    baseline_mtime = os.path.getmtime(baseline) if baseline else 0

    if not _open_message_dropdown_menu(driver, bubble):
        print("⚠️ [VOICE] Could not open message dropdown menu.")
        return None

    if not _click_download_in_context_menu(driver):
        print("⚠️ [VOICE] Download item not found in message menu.")
        try:
            ActionChains(driver).send_keys(Keys.ESCAPE).perform()
        except Exception:
            pass
        return None

    for _ in range(30):
        time.sleep(0.5)
        path = _pick_newest_voice_download(min_mtime=baseline_mtime)
        if not path:
            continue
        try:
            with open(path, "rb") as f:
                raw = f.read()
            if is_valid_voice_bytes(raw):
                print(f"✅ [VOICE] Got opus from Downloads: {path} ({len(raw)} bytes)")
                return raw
        except OSError as exc:
            print(f"⚠️ [VOICE] Read download failed: {exc}")

    print("⚠️ [VOICE] Menu Download did not produce a new .opus within 15s.")
    return None


def download_voice_from_bubble(
    driver,
    bubble,
    contact_name: str,
    message_data_id: str = "",
) -> Optional[str]:
    install_voice_capture_hooks(driver)
    container = _relocate_voice_container(driver, bubble, message_data_id)
    if container is None:
        print("⚠️ [VOICE] No message container for voice download.")
        return None

    os.makedirs(WA_AUDIO_DIR, exist_ok=True)

    try:
        try:
            driver.execute_script(
                "var c=document.querySelector('[data-testid=\"btn-closer-drawer\"]'); if(c) c.click();"
            )
            time.sleep(0.6)
        except Exception:
            pass
        driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center', inline: 'center'});",
            container,
        )
        time.sleep(0.8)
        data_id = _resolve_data_id(container, message_data_id)
        id_preview = data_id[:48] if data_id else "(none)"

        print(f"🎤 [VOICE] Target WA_Audio download data_id={id_preview!r}...")

        raw = _try_download_voice_via_button(driver, container)
        if raw is not None:
            saved = _save_voice_to_wa_audio(raw, data_id)
            if saved:
                return saved

        print("⚠️ [VOICE] message menu Download failed — trying whatsapp-store...")

        raw = _try_download_b64(
            driver, "whatsapp-store", DOWNLOAD_VOICE_NO_PLAY_JS, data_id, container
        )
        if raw is not None and is_valid_voice_bytes(raw):
            saved = _save_voice_to_wa_audio(raw, data_id)
            if saved:
                return saved
            print("⚠️ [VOICE] whatsapp-store bytes invalid — trying capture-hooked-blob...")
        elif raw is not None:
            header = raw[:16].hex() if raw else "empty"
            print(
                f"⚠️ [VOICE] whatsapp-store returned non-opus ({len(raw)} bytes, header={header}) "
                "— trying capture-hooked-blob..."
            )
        else:
            print("⚠️ [VOICE] whatsapp-store returned no data — trying capture-hooked-blob...")

        raw = _try_capture_hooked_blob(driver, container)
        if raw is not None:
            saved = _save_voice_to_wa_audio(raw, data_id)
            if saved:
                return saved
            header = raw[:16].hex() if raw else "empty"
            print(
                f"⚠️ [VOICE] capture-hooked-blob non-opus ({len(raw)} bytes, header={header}) "
                "— trying capture-ogg-blob..."
            )

        raw = _try_capture_ogg_audio(driver, container)
        if raw is not None:
            saved = _save_voice_to_wa_audio(raw, data_id)
            if saved:
                return saved

        hooked_count = driver.execute_script(
            "return (window.__openclawVoiceUrls || []).length;"
        )
        print(
            f"⚠️ [VOICE] Download failed — no file in {WA_AUDIO_DIR} "
            f"(hooked_urls={hooked_count}, data_id={data_id[:32]!r})"
        )
        return None
    except Exception as exc:
        print(f"❌ [VOICE] Download failed: {exc}")
        return None


def _whisper_language_label() -> str:
    lang = (WHISPER_LANGUAGE or "auto").strip().lower()
    return lang if lang not in ("", "auto", "detect") else "auto"


def _whisper_transcribe_opts() -> Dict[str, Any]:
    opts: Dict[str, Any] = {
        "fp16": False,
        "condition_on_previous_text": False,
        "initial_prompt": WHISPER_PROMPT,
    }
    lang = (WHISPER_LANGUAGE or "").strip().lower()
    if lang and lang not in ("auto", "detect"):
        opts["language"] = lang
    return opts


def normalize_transcript_via_copilot(raw_transcript: str, detected_language: str = "") -> str:
    """Optional Copilot pass to turn mixed EN/MY/ZH speech into clear RFQ English."""
    text = str(raw_transcript or "").strip()
    if not text:
        return ""
    lang_hint = f" Detected speech language: {detected_language}." if detected_language else ""
    prompt = (
        "Normalize this WhatsApp voice transcript from a Malaysian industrial-parts customer. "
        "The speaker may mix English, Malay, and Mandarin. "
        "Return one clear English RFQ sentence. Keep part numbers, quantities, and units exact."
        f"{lang_hint}\n\nTranscript:\n{text}"
    )
    client = OpenAI(
        base_url=COPILOT_BASE_URL,
        api_key=os.getenv("COPILOT_API_KEY", "local-copilot-proxy"),
        timeout=60.0,
        max_retries=1,
    )
    try:
        response = client.chat.completions.create(
            model=COPILOT_MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
        )
        normalized = str(response.choices[0].message.content or "").strip()
        if normalized:
            print(f"🎤 [VOICE] Copilot normalized transcript: {normalized[:200]}")
            return normalized
    except Exception as exc:
        print(f"⚠️ [VOICE] Copilot transcript normalization skipped: {exc}")
    return text


def _opus_path_for_transcription(preferred_path: str = "") -> str:
    """Prefer saved .opus (curl-equivalent) over ffmpeg .wav sidecar."""
    for path in (preferred_path, VOICE_LATEST_OPUS):
        if path and os.path.exists(path) and path.lower().endswith((".opus", ".ogg")):
            return path
    if preferred_path and os.path.exists(preferred_path):
        return preferred_path
    return VOICE_LATEST_OPUS if os.path.exists(VOICE_LATEST_OPUS) else (preferred_path or "")


def _file_is_recent(path: str, max_age_seconds: int = 300) -> bool:
    try:
        return (time.time() - os.path.getmtime(path)) <= max_age_seconds
    except OSError:
        return False


def transcribe_audio_via_copilot(audio_path: str) -> Tuple[str, str]:
    """Transcribe through the local Copilot proxy (POST /v1/audio/transcriptions)."""
    upload_path = _opus_path_for_transcription(audio_path)
    if not upload_path or not os.path.exists(upload_path):
        print(f"⚠️ [VOICE] No audio file to upload (preferred={audio_path!r}).")
        return "", ""

    url = f"{COPILOT_BASE_URL.rstrip('/')}/audio/transcriptions"
    size = os.path.getsize(upload_path)
    ext = os.path.splitext(upload_path)[1] or ".opus"
    print(
        f"🎤 [VOICE] POST {url} model={COPILOT_WHISPER_MODEL} "
        f"language={_whisper_language_label()} bytes={size} format={ext} "
        f"file={os.path.basename(upload_path)}"
    )
    upload_name = os.path.basename(upload_path)
    data = {"model": COPILOT_WHISPER_MODEL}
    lang = (WHISPER_LANGUAGE or "").strip().lower()
    if lang and lang not in ("auto", "detect", ""):
        data["language"] = lang
    mime = "audio/ogg" if ext.lower() in (".opus", ".ogg") else "audio/wav"
    try:
        with open(upload_path, "rb") as audio_file:
            response = requests.post(
                url,
                files={
                    "file": (
                        upload_name,
                        audio_file,
                        mime,
                    )
                },
                data=data,
                timeout=180,
            )
        response.raise_for_status()
        payload = response.json()
        text = str(payload.get("text") or "").strip()
        detected = str(payload.get("language") or "").strip().lower()
        if detected:
            print(f"🎤 [VOICE] Detected language: {detected}")
        if text:
            print(f"🎤 [VOICE] Copilot transcript: {text[:200]}")
        else:
            print("⚠️ [VOICE] Copilot transcription returned empty text.")
        return text, detected
    except requests.exceptions.HTTPError as exc:
        body = ""
        if exc.response is not None:
            body = (exc.response.text or "")[:400]
        print(f"⚠️ [VOICE] Copilot transcription HTTP error ({url}): {exc} {body}")
        return "", ""
    except Exception as exc:
        print(f"⚠️ [VOICE] Copilot transcription failed ({url}): {exc}")
        return "", ""


def _clear_workspace_dir(directory: str, label: str) -> None:
    """Remove all files in a WA_* workspace after processing."""
    if not os.path.isdir(directory):
        return
    removed = 0
    for name in os.listdir(directory):
        path = os.path.join(directory, name)
        if os.path.isfile(path):
            try:
                os.remove(path)
                removed += 1
            except OSError:
                pass
    if removed:
        print(f"🧹 [{label}] Cleared {removed} file(s) from {directory}")


def clear_wa_audio_workspace():
    """Remove saved voice files after processing so the next message starts clean."""
    _clear_workspace_dir(WA_AUDIO_DIR, "VOICE")


def clear_wa_image_workspace():
    """Remove saved image screenshots after Copilot analysis."""
    _clear_workspace_dir(WA_IMAGE_DIR, "IMAGE")


def clear_wa_files_workspace():
    """Remove downloaded document files after Copilot analysis."""
    _clear_workspace_dir(WA_FILES_DIR, "FILES")


def clear_wa_attachment_workspace():
    """Clear all temporary WhatsApp attachment workspaces (audio, image, files)."""
    clear_wa_audio_workspace()
    clear_wa_image_workspace()
    clear_wa_files_workspace()


def save_wa_image_manifest(image_path: str, message_data_id: str = "") -> None:
    """Track the current message image like latest.opus for voice."""
    if not image_path or not os.path.exists(image_path):
        return
    os.makedirs(WA_IMAGE_DIR, exist_ok=True)
    try:
        shutil.copy2(image_path, IMAGE_LATEST_PATH)
    except OSError as exc:
        print(f"⚠️ [IMAGE] Could not update {IMAGE_LATEST_PATH}: {exc}")
        return
    manifest_path = os.path.join(WA_IMAGE_DIR, "latest.json")
    try:
        with open(manifest_path, "w", encoding="utf-8") as f:
            json.dump(
                {
                    "saved_at": datetime.datetime.now().isoformat(timespec="seconds"),
                    "data_id": message_data_id,
                    "named_path": image_path,
                    "bytes": os.path.getsize(image_path),
                },
                f,
                indent=2,
            )
    except Exception:
        pass
    print(f"✅ [IMAGE] Updated {IMAGE_LATEST_PATH}")


def _read_voice_manifest() -> dict:
    manifest_path = os.path.join(WA_AUDIO_DIR, "latest.json")
    if not os.path.exists(manifest_path):
        return {}
    try:
        with open(manifest_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def ensure_voice_transcript(
    voice_path: str = "",
    message_data_id: str = "",
    max_age_seconds: int = 120,
) -> str:
    """
    Transcribe a voice note only when latest.opus matches this message data_id.
    Never reuse a previous customer's recording.
    """
    expected_id = str(message_data_id or "").strip()
    manifest = _read_voice_manifest()
    manifest_id = str(manifest.get("data_id") or "").strip()

    if not os.path.exists(VOICE_LATEST_OPUS):
        print("⚠️ [VOICE] No latest.opus — download may have failed for this bubble.")
        return ""

    if expected_id:
        if not manifest_id:
            print(
                f"⚠️ [VOICE] Refusing stale opus — no manifest for data_id={expected_id[:28]!r}"
            )
            return ""
        if expected_id not in manifest_id and manifest_id not in expected_id:
            print(
                f"⚠️ [VOICE] Refusing stale opus — manifest={manifest_id[:28]!r} "
                f"≠ message={expected_id[:28]!r}"
            )
            return ""
    elif not _file_is_recent(VOICE_LATEST_OPUS, max_age_seconds):
        print("⚠️ [VOICE] latest.opus too old and no data_id — skip transcript")
        return ""

    upload_path = voice_path if voice_path and os.path.exists(voice_path) else VOICE_LATEST_OPUS
    if not _file_is_recent(upload_path, max_age_seconds):
        print(f"⚠️ [VOICE] Audio file not fresh enough: {upload_path}")
        return ""

    text = transcribe_audio(upload_path)
    return text


def transcribe_audio_via_copilot_chat(audio_path: str, caption: str = "") -> str:
    """Not used — the Copilot bridge only forwards images on /v1/chat/completions."""
    return ""


def transcribe_audio(audio_path: str, caption: str = "") -> str:
    """Upload saved .opus to Copilot /v1/audio/transcriptions — identical to curl."""
    if not audio_path or not os.path.exists(audio_path):
        print("⚠️ [VOICE] No audio file to transcribe.")
        return ""

    text, detected = transcribe_audio_via_copilot(audio_path)
    if text and WHISPER_NORMALIZE:
        text = normalize_transcript_via_copilot(text, detected_language=detected)
    return text


def download_document_from_bubble(
    driver,
    bubble,
    contact_name: str,
    filename: str = "",
) -> Optional[str]:
    element = _find_document_element(bubble)
    if element is None:
        return None

    os.makedirs(DOC_CAPTURE_DIR, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", filename or "document")[:80]
    if not safe_name or safe_name == "document":
        safe_name = "document.bin"
    out_path = os.path.join(
        DOC_CAPTURE_DIR,
        f"{_timestamp_slug()}_{_safe_contact(contact_name)}_{safe_name}",
    )

    try:
        driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center', inline: 'center'});",
            element,
        )
        time.sleep(0.5)

        inline_download = _find_download_in_bubble(bubble)
        if inline_download is not None:
            inline_download.click()
            time.sleep(4)
            downloaded = _pick_newest_download(safe_name)
            if downloaded:
                os.replace(downloaded, out_path)
                print(f"📄 [DOC] Saved document (inline download) → {out_path}")
                return out_path

        element.click()
        time.sleep(2)

        download_selectors = [
            '[data-testid="download"]',
            '[data-icon="download"]',
            'span[data-icon="download"]',
            '[aria-label="Download"]',
            '[title="Download"]',
            '[data-testid="media-viewer-download"]',
        ]
        for selector in download_selectors:
            try:
                buttons = driver.find_elements(By.CSS_SELECTOR, selector)
                for btn in buttons:
                    if btn.is_displayed():
                        btn.click()
                        time.sleep(3)
                        break
            except Exception:
                continue

        downloaded = _pick_newest_download(safe_name)
        if downloaded:
            os.replace(downloaded, out_path)
            print(f"📄 [DOC] Saved document → {WA_FILES_DIR}/{os.path.basename(out_path)}")
            return out_path

        b64 = _execute_async_js(driver, BLOB_TO_BASE64_JS, element)
        if b64:
            with open(out_path, "wb") as f:
                f.write(base64.b64decode(b64))
            print(f"📄 [DOC] Saved document via blob: {out_path}")
            return out_path

        print("⚠️ [DOC] Could not download document attachment.")
        preview_path = capture_document_preview_image(driver, bubble, contact_name)
        if preview_path:
            return preview_path
        return None
    except Exception as exc:
        print(f"❌ [DOC] Document download failed: {exc}")
        return None
    finally:
        try:
            close_buttons = driver.find_elements(By.CSS_SELECTOR, '[data-testid="x"]')
            for btn in close_buttons:
                if btn.is_displayed():
                    btn.click()
                    break
        except Exception:
            pass


MIN_WA_IMAGE_BYTES = int(os.getenv("OPENCLAW_MIN_WA_IMAGE_BYTES", "8000"))
MIN_WA_IMAGE_DISPLAY_PX = int(os.getenv("OPENCLAW_MIN_WA_IMAGE_DISPLAY_PX", "400"))
MIN_WA_IMAGE_NATURAL_PX = int(os.getenv("OPENCLAW_MIN_WA_IMAGE_NATURAL_PX", "250"))


def detect_image_format(data: bytes) -> Tuple[Optional[str], str]:
    """Return (mime_type, file_extension) from image bytes."""
    if len(data) >= 8 and data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png", ".png"
    if len(data) >= 2 and data[:2] == b"\xff\xd8":
        return "image/jpeg", ".jpg"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", ".webp"
    return None, ""


def read_image_dimensions(path: str) -> Optional[Tuple[int, int]]:
    """Return (width, height) from PNG/JPEG/WebP header without Pillow."""
    if not path or not os.path.exists(path):
        return None
    try:
        with open(path, "rb") as handle:
            data = handle.read(256 * 1024)
    except OSError:
        return None
    if len(data) >= 24 and data[:8] == b"\x89PNG\r\n\x1a\n":
        w = int.from_bytes(data[16:20], "big")
        h = int.from_bytes(data[20:24], "big")
        if w > 0 and h > 0:
            return w, h
    if len(data) >= 2 and data[:2] == b"\xff\xd8":
        idx = 2
        while idx + 9 < len(data):
            if data[idx] != 0xFF:
                idx += 1
                continue
            marker = data[idx + 1]
            if marker in (0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF):
                h = int.from_bytes(data[idx + 5:idx + 7], "big")
                w = int.from_bytes(data[idx + 7:idx + 9], "big")
                if w > 0 and h > 0:
                    return w, h
                break
            if marker in (0xD0, 0xD1, 0xD2, 0xD3, 0xD4, 0xD5, 0xD6, 0xD7, 0xD8, 0xD9, 0x01):
                idx += 4
                continue
            if idx + 3 >= len(data):
                break
            seg_len = int.from_bytes(data[idx + 2:idx + 4], "big")
            idx += 2 + max(seg_len, 2)
    if len(data) >= 30 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        chunk = data[12:16]
        if chunk == b"VP8 " and len(data) >= 30:
            w = int.from_bytes(data[26:28], "little") & 0x3FFF
            h = int.from_bytes(data[28:30], "little") & 0x3FFF
            if w > 0 and h > 0:
                return w, h
        if chunk == b"VP8L" and len(data) >= 25:
            bits = int.from_bytes(data[21:25], "little")
            w = (bits & 0x3FFF) + 1
            h = ((bits >> 14) & 0x3FFF) + 1
            if w > 0 and h > 0:
                return w, h
    return None


def validate_image_file(path: str, min_bytes: int = None) -> Tuple[bool, str]:
    """Reject corrupt placeholders / tiny thumbnails before Copilot vision."""
    min_bytes = MIN_WA_IMAGE_BYTES if min_bytes is None else min_bytes
    if not path or not os.path.exists(path):
        return False, "file missing"
    size = os.path.getsize(path)
    if size < min_bytes:
        return False, f"too small ({size} bytes, need >= {min_bytes})"
    try:
        with open(path, "rb") as handle:
            head = handle.read(16)
    except OSError as exc:
        return False, str(exc)
    mime, _ext = detect_image_format(head)
    if not mime:
        return False, "not a valid PNG/JPEG/WebP image"
    return True, mime


def save_validated_image_bytes(data: bytes, dest_path: str, min_bytes: int = None) -> Optional[str]:
    """Write image bytes with correct extension; return path or None if invalid."""
    min_bytes = MIN_WA_IMAGE_BYTES if min_bytes is None else min_bytes
    mime, ext = detect_image_format(data)
    if not mime or len(data) < min_bytes:
        return None
    base, _ = os.path.splitext(dest_path)
    final_path = f"{base}{ext}"
    with open(final_path, "wb") as handle:
        handle.write(data)
    ok, reason = validate_image_file(final_path, min_bytes=min_bytes)
    if not ok:
        try:
            os.remove(final_path)
        except OSError:
            pass
        print(f"⚠️ [IMAGE] Rejected saved image: {reason}")
        return None
    return final_path


def pick_newest_image_download(baseline_mtime: float = 0, min_bytes: int = None) -> Optional[str]:
    """Find a freshly downloaded image in ~/Downloads (Save Image As / viewer download)."""
    download_dirs = [
        os.path.expanduser("~/Downloads"),
        WA_IMAGE_DIR,
    ]
    image_exts = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")
    candidates = []
    for directory in download_dirs:
        if not os.path.isdir(directory):
            continue
        for name in os.listdir(directory):
            path = os.path.join(directory, name)
            if not os.path.isfile(path):
                continue
            if not name.lower().endswith(image_exts):
                continue
            mtime = os.path.getmtime(path)
            if baseline_mtime and mtime <= baseline_mtime:
                continue
            candidates.append(path)
    if not candidates:
        return None
    candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    min_bytes = MIN_WA_IMAGE_BYTES if min_bytes is None else min_bytes
    for path in candidates:
        if time.time() - os.path.getmtime(path) > 120:
            continue
        ok, _reason = validate_image_file(path, min_bytes=min_bytes)
        if ok:
            return path
    return None


def _pick_newest_download(preferred_name: str) -> Optional[str]:
    download_dirs = [
        os.path.expanduser("~/Downloads"),
        WA_FILES_DIR,
    ]
    candidates = []
    for directory in download_dirs:
        if not os.path.isdir(directory):
            continue
        for name in os.listdir(directory):
            path = os.path.join(directory, name)
            if not os.path.isfile(path):
                continue
            if preferred_name and preferred_name.lower() in name.lower():
                candidates.append(path)
            elif name.lower().endswith((
                ".pdf", ".xlsx", ".xls", ".docx", ".doc", ".pptx", ".ppt", ".csv"
            )):
                candidates.append(path)
    if not candidates:
        return None
    candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    newest = candidates[0]
    if time.time() - os.path.getmtime(newest) > 120:
        return None
    return newest


def extract_text_from_document(file_path: str) -> str:
    if not file_path or not os.path.exists(file_path):
        return ""

    lower = file_path.lower()
    try:
        if lower.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages[:20]:
                pages.append(page.extract_text() or "")
            return "\n".join(pages).strip()

        if lower.endswith((".xlsx", ".xls", ".xlsm")):
            import pandas as pd

            frames = pd.read_excel(file_path, sheet_name=None, header=None)
            chunks = []
            for _name, df in frames.items():
                chunks.append(df.fillna("").astype(str).to_string(index=False, header=False))
            return "\n\n".join(chunks).strip()

        if lower.endswith((".docx",)):
            from docx import Document

            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()

        if lower.endswith((".csv",)):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(50000).strip()

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(10000).strip()
    except Exception as exc:
        print(f"⚠️ [DOC] Text extraction failed for {file_path}: {exc}")
        return ""


def extract_items_from_document_text(document_text: str, file_path: str = "") -> List[Dict[str, Any]]:
    text = str(document_text or "").strip()
    if not text:
        return []

    client = OpenAI(
        base_url=COPILOT_BASE_URL,
        api_key=os.getenv("COPILOT_API_KEY", "local-copilot-proxy"),
        timeout=45.0,
        max_retries=1,
    )
    prompt = (
        "Extract industrial parts / PO line items from this document text.\n"
        "Return STRICT JSON array: "
        '[{"part_no":"...", "qty":1, "brand":"UNKNOWN"}]\n'
        "Use positive integer qty. Do not invent part numbers.\n\n"
        f"Filename: {os.path.basename(file_path or '')}\n"
        f"Document text:\n{text[:12000]}"
    )
    try:
        response = client.chat.completions.create(
            model=COPILOT_MODEL,
            messages=[
                {"role": "system", "content": "You extract structured PO/RFQ line items."},
                {"role": "user", "content": prompt},
            ],
        )
        raw = (response.choices[0].message.content or "").strip()
        if raw.startswith("```"):
            raw = "\n".join(raw.splitlines()[1:-1]).strip()
        parsed = json.loads(raw)
        if not isinstance(parsed, list):
            return []
        items = []
        for row in parsed:
            if not isinstance(row, dict):
                continue
            part_no = str(row.get("part_no") or "").strip().upper()
            try:
                qty = int(row.get("qty") or 1)
            except (TypeError, ValueError):
                qty = 1
            if part_no and qty > 0:
                items.append({
                    "part_no": part_no,
                    "qty": qty,
                    "brand": str(row.get("brand") or "UNKNOWN").strip().upper(),
                })
        return items
    except Exception as exc:
        print(f"⚠️ [DOC] Copilot line-item extraction failed: {exc}")
        return []


def enrich_message_from_attachments(
    driver,
    bubble,
    contact_name: str,
    latest_message: str,
    media_info,
    message_data_id: str = "",
) -> Dict[str, Any]:
    """
    Return enriched text plus optional extracted document items and local file paths.
    """
    result = {
        "text": latest_message or "",
        "voice_path": "",
        "document_path": "",
        "document_text": "",
        "document_items": [],
        "transcript": "",
        "preview_image_path": "",
    }

    media_type = getattr(media_info, "media_type", "text")
    filename = getattr(media_info, "filename", "") or ""

    if media_type == "voice" or getattr(media_info, "has_voice", False):
        print(f"🎤 [VOICE] Processing voice attachment (caption={latest_message!r})...")
        if bubble is None:
            print("⚠️ [VOICE] Bubble element is None — cannot download.")
        else:
            voice_path = download_voice_from_bubble(
                driver, bubble, contact_name, message_data_id=message_data_id
            )
            manifest = _read_voice_manifest()
            manifest_id = str(manifest.get("data_id") or "").strip()
            expected_id = str(message_data_id or "").strip()
            has_fresh_opus = (
                os.path.exists(VOICE_LATEST_OPUS)
                and manifest_id
                and (
                    not expected_id
                    or expected_id in manifest_id
                    or manifest_id in expected_id
                )
            )
            if voice_path and has_fresh_opus:
                result["voice_path"] = voice_path
                transcript = ensure_voice_transcript(
                    voice_path,
                    message_data_id=message_data_id,
                )
                result["transcript"] = transcript
                if transcript:
                    prefix = f"[Voice transcript]\n{transcript}"
                    result["text"] = f"{prefix}\n\n{latest_message}".strip() if latest_message else prefix
                elif latest_message:
                    result["text"] = latest_message
            elif latest_message:
                result["text"] = latest_message
            elif not has_fresh_opus:
                print("⚠️ [VOICE] Download failed — no valid opus saved for this message.")

    doc_types = ("pdf", "office_word", "office_excel", "office_powerpoint", "document")
    is_doc = (
        media_type in doc_types
        or getattr(media_info, "has_document", False)
        or filename.lower().endswith((".pdf", ".xlsx", ".xls", ".doc", ".docx"))
    )
    if is_doc:
        print(f"📄 [DOC] Processing document attachment: {filename or media_type}")
        doc_path = download_document_from_bubble(
            driver, bubble, contact_name, filename=filename
        )
        if doc_path:
            result["document_path"] = doc_path
            if doc_path.lower().endswith(".pdf"):
                doc_text = extract_text_from_document(doc_path)
            elif doc_path.lower().endswith(".png"):
                result["preview_image_path"] = doc_path
                doc_text = ""
            else:
                doc_text = extract_text_from_document(doc_path)

            if not doc_text and doc_path.lower().endswith(".png"):
                result["preview_image_path"] = doc_path
            elif doc_text:
                result["document_text"] = doc_text
                header = f"[Document: {os.path.basename(doc_path)}]\n{doc_text[:4000]}"
                if "PURCHASE ORDER" in doc_text.upper() or filename.lower().endswith(".pdf"):
                    header = f"[Purchase Order PDF: {filename or os.path.basename(doc_path)}]\n{doc_text[:4000]}"
                result["text"] = f"{header}\n\n{latest_message}".strip() if latest_message else header
                result["document_items"] = extract_items_from_document_text(doc_text, doc_path)

            if not doc_text and not result["document_items"]:
                preview = capture_document_preview_image(driver, bubble, contact_name)
                if preview:
                    result["preview_image_path"] = preview
                    result["text"] = (
                        f"[Purchase Order PDF preview: {filename}]\n"
                        f"(Download failed — using preview screenshot for extraction)\n\n"
                        f"{latest_message or ''}"
                    ).strip()

    return result
